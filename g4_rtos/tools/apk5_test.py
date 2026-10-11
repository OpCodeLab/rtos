"""APK5 automated test runner: sends frame patterns to the STM32 and writes a PASS/FAIL report.

Each test case matches an ID of docs/APK5_Requirements_ASPICE.xlsx (sheet "Test Specification").

Usage:
    pip install pyserial                 (openpyxl too, only for --xlsx)
    python apk5_test.py --port COM5                      run all tests on the board
    python apk5_test.py --port COM5 --only TC-01 TC-08   run a subset
    python apk5_test.py --port COM5 --interactive        also run the button test (operator presses B1)
    python apk5_test.py --sim                            run against a built-in board simulator
    python apk5_test.py --port COM5 --xlsx ../docs/APK5_Requirements_ASPICE.xlsx --tester "M. Hanzouti"
                                                         also write Result / Date / Tester into the workbook

Reports are written to tools/reports/apk5_report_<date>_<time>.html and .csv
"""
import argparse
import csv
import datetime
import html
import os
import queue
import re
import sys
import threading
import time

SOF = 0xAA
RSP_FLAG = 0x80
FRAME_TIMEOUT_S = 0.05

CMD_PING = 0x01
CMD_LED_SET = 0x02
CMD_LED_TOGGLE = 0x03
CMD_GET_BUTTON = 0x04

ST_OK = 0x00
ST_BAD_CMD = 0xE1
ST_BAD_DATA = 0xE2

ANSI_RE = re.compile(r"\x1b\[[0-9;]*m")


def hexs(data):
    return " ".join(f"{b:02X}" for b in data)


# ---------------------------------------------------------------- link ---
class Link:
    """Serial link with a background reader running the SOF/CMD/DATA state machine."""

    def __init__(self, port):
        self.port = port
        self.events = queue.Queue()     # ("frame", bytes3, t) or ("text", str, t)
        self.running = True
        self.thread = threading.Thread(target=self._reader, daemon=True)
        self.thread.start()

    def _reader(self):
        state, cmd, last, line = 0, 0, 0.0, bytearray()
        while self.running:
            chunk = self.port.read(64)
            for byte in chunk:
                now = time.perf_counter()
                if state and now - last > FRAME_TIMEOUT_S:
                    state = 0
                last = now
                if state == 0:
                    if byte == SOF:
                        state = 1
                    elif byte == 0x0A:
                        text = ANSI_RE.sub("", line.decode(errors="replace")).strip()
                        if text:
                            self.events.put(("text", text, now))
                        line.clear()
                    elif byte != 0x0D:
                        line.append(byte)
                elif state == 1:
                    cmd, state = byte, 2
                else:
                    self.events.put(("frame", bytes([SOF, cmd, byte]), now))
                    state = 0

    def close(self):
        self.running = False
        self.thread.join(1)
        self.port.close()

    def flush(self, settle=0.1):
        time.sleep(settle)
        while not self.events.empty():
            self.events.get_nowait()

    def write(self, data):
        self.port.write(bytes(data))
        return time.perf_counter()

    def collect(self, timeout, max_frames=None):
        """Return (frames, texts) received within timeout seconds."""
        frames, texts = [], []
        end = time.perf_counter() + timeout
        while True:
            left = end - time.perf_counter()
            if left <= 0 or (max_frames and len(frames) >= max_frames):
                break
            try:
                kind, value, t = self.events.get(timeout=left)
            except queue.Empty:
                break
            (frames if kind == "frame" else texts).append((value, t))
        return frames, texts

    def transact(self, cmd, data, timeout=0.3):
        """Send one frame, return (response bytes or None, round-trip seconds, log texts)."""
        t0 = self.write([SOF, cmd, data])
        frames, texts = self.collect(timeout, max_frames=1)
        if not frames:
            return None, None, [t for t, _ in texts]
        rsp, t1 = frames[0]
        more, extra = self.collect(0.02)          # trailing log lines
        return rsp, t1 - t0, [t for t, _ in texts + extra]


# ---------------------------------------------------------- simulator ---
class SimBoard:
    """Pretends to be the STM32 running APK5 (same protocol, 8-frame queue, log lines)."""

    def __init__(self):
        self.rx = queue.Queue()
        self.out = bytearray()
        self.lock = threading.Lock()
        self.leds = 0
        self.button = 0
        self.state, self.cmd, self.last = 0, 0, 0.0
        threading.Thread(target=self._task, daemon=True).start()

    # pyserial-like API
    def write(self, data):
        for byte in data:
            self._isr(byte)

    def read(self, n):
        time.sleep(0.002)
        with self.lock:
            chunk, self.out[:] = bytes(self.out[:n]), self.out[n:]
        return chunk

    def close(self):
        pass

    def _isr(self, byte):
        now = time.perf_counter()
        if self.state and now - self.last > FRAME_TIMEOUT_S:
            self.state = 0
        self.last = now
        if self.state == 0:
            self.state = 1 if byte == SOF else 0
        elif self.state == 1:
            self.cmd, self.state = byte, 2
        else:
            if self.rx.qsize() < 8:
                self.rx.put((self.cmd, byte))
            self.state = 0

    def _task(self):
        while True:
            cmd, data = self.rx.get()
            time.sleep(0.003)                         # printf of the trace line
            if cmd == CMD_PING:
                rsp = data
            elif cmd == CMD_LED_SET:
                self.leds, rsp = data & 0x0F, ST_OK
            elif cmd == CMD_LED_TOGGLE:
                if data > 3:
                    rsp = ST_BAD_DATA
                else:
                    self.leds ^= 1 << data
                    rsp = ST_OK
            elif cmd == CMD_GET_BUTTON:
                rsp = self.button
            else:
                rsp = ST_BAD_CMD
            log = f"\x1b[36mRX frame CMD=0x{cmd:02X} DATA=0x{data:02X}\x1b[0m\r\n".encode()
            with self.lock:
                self.out += log + bytes([SOF, cmd | RSP_FLAG, rsp])


# -------------------------------------------------------------- tests ---
TESTS = []


def testcase(tid, title, verifies, interactive=False):
    def deco(fn):
        TESTS.append({"id": tid, "title": title, "verifies": verifies, "fn": fn, "interactive": interactive})
        return fn
    return deco


class Ctx:
    """Passed to each test: the link plus a check() that records every step."""

    def __init__(self, link, args):
        self.link, self.args = link, args
        self.steps = []
        self.rtts = []

    def check(self, step, sent, expected, received, ok, note=""):
        self.steps.append({"step": step, "sent": sent, "expected": expected,
                           "received": received, "ok": bool(ok), "note": note})
        return ok

    def expect(self, step, cmd, data, exp_data, timeout=0.3):
        """Send one frame and check the response is AA (cmd|0x80) exp_data."""
        rsp, rtt, _ = self.link.transact(cmd, data, timeout)
        exp = bytes([SOF, (cmd | RSP_FLAG) & 0xFF, exp_data])
        if rtt is not None:
            self.rtts.append(rtt)
        note = f"{rtt * 1000:.1f} ms" if rtt is not None else "timeout"
        return self.check(step, hexs([SOF, cmd, data]), hexs(exp),
                          hexs(rsp) if rsp else "-", rsp == exp, note)


@testcase("TC-01", "PING echo, data patterns", "SW-REQ-001, 002, 004, 007, 008, 014, 015, 016")
def tc_ping(c):
    for d in (0x00, 0x01, 0x55, 0x7F, 0x80, 0xAA, 0xFE, 0xFF):
        c.expect(f"PING 0x{d:02X}", CMD_PING, d, d)
    c.link.flush(0)
    _, _, texts = c.link.transact(CMD_PING, 0x33)
    trace = any("RX frame CMD=0x01 DATA=0x33" in t for t in texts)
    c.check("Trace line printed", "AA 01 33", "log 'RX frame CMD=0x01 DATA=0x33'",
            "; ".join(texts) or "-", trace)


@testcase("TC-02", "LED_SET, all 16 masks", "SW-REQ-009, 027")
def tc_led_set(c):
    for mask in range(16):
        c.expect(f"LED_SET 0x{mask:X}", CMD_LED_SET, mask, ST_OK)
        time.sleep(c.args.led_delay)
    c.expect("LED_SET 0x00 (all off)", CMD_LED_SET, 0x00, ST_OK)


@testcase("TC-03", "LED_TOGGLE valid index", "SW-REQ-010")
def tc_toggle(c):
    for _ in range(2):                       # twice: LEDs end in their initial state
        for i in range(4):
            c.expect(f"TOGGLE {i}", CMD_LED_TOGGLE, i, ST_OK)
            time.sleep(c.args.led_delay)


@testcase("TC-04", "LED_TOGGLE out of range", "SW-REQ-013")
def tc_toggle_bad(c):
    for d in (4, 7, 0x10, 0x80, 0xFF):
        c.expect(f"TOGGLE {d} rejected", CMD_LED_TOGGLE, d, ST_BAD_DATA)


@testcase("TC-05", "Unknown commands", "SW-REQ-012, 014, 028")
def tc_unknown(c):
    for cmd in (0x00, 0x05, 0x42, 0x7F, 0x81):
        c.expect(f"CMD 0x{cmd:02X} rejected", cmd, 0x00, ST_BAD_CMD)


@testcase("TC-06", "GET_BUTTON", "SW-REQ-011")
def tc_button(c):
    rsp, _, _ = c.link.transact(CMD_GET_BUTTON, 0)
    c.check("Response format", "AA 04 00", "AA 84 00|01", hexs(rsp) if rsp else "-",
            rsp is not None and rsp[1] == 0x84 and rsp[2] in (0, 1))
    if not c.args.interactive:
        c.check("Pressed / released", "-", "-", "-", True, "skipped: run with --interactive")
        return
    input("    >>> HOLD the blue user button B1, then press Enter ")
    c.expect("Button pressed", CMD_GET_BUTTON, 0, 1)
    input("    >>> RELEASE the button, then press Enter ")
    c.expect("Button released", CMD_GET_BUTTON, 0, 0)


@testcase("TC-07", "Leading garbage before SOF", "SW-REQ-003")
def tc_garbage(c):
    patterns = [
        [0x12, 0x34],
        [0x00, 0xFF, 0x55],
        list(b"hello\r\n"),
    ]
    for g in patterns:
        c.link.flush(0.05)
        c.link.write(g + [SOF, CMD_PING, 0x10])
        frames, _ = c.link.collect(0.3)
        got = [hexs(f) for f, _ in frames]
        c.check(f"Garbage {hexs(g)}", hexs(g + [SOF, CMD_PING, 0x10]), "AA 81 10 (once)",
                ", ".join(got) or "-", got == ["AA 81 10"])


@testcase("TC-08", "Inter-byte timeout drops partial frame", "SW-REQ-005")
def tc_timeout(c):
    for partial in ([SOF], [SOF, CMD_PING]):
        c.link.flush(0.05)
        c.link.write(partial)
        time.sleep(0.2)                        # > 50 ms firmware timeout
        c.link.write([SOF, CMD_PING, 0x20])
        frames, _ = c.link.collect(0.3)
        got = [hexs(f) for f, _ in frames]
        c.check(f"Partial {hexs(partial)} + 200 ms", f"{hexs(partial)} .. AA 01 20", "AA 81 20 (once)",
                ", ".join(got) or "-", got == ["AA 81 20"])
    # bytes split by less than the timeout must still form one frame
    c.link.flush(0.05)
    for b in (SOF, CMD_PING, 0x21):
        c.link.write([b])
        time.sleep(0.01)
    frames, _ = c.link.collect(0.3)
    got = [hexs(f) for f, _ in frames]
    c.check("Bytes 10 ms apart", "AA . 01 . 21", "AA 81 21", ", ".join(got) or "-", got == ["AA 81 21"])


@testcase("TC-09", "Burst of frames (queue depth 8)", "SW-REQ-007, 014, 020")
def tc_burst(c):
    for count, must_all in ((8, True), (20, False)):
        c.link.flush(0.2)
        data = []
        for i in range(count):
            data += [SOF, CMD_PING, i]
        c.link.write(data)
        frames, _ = c.link.collect(1.0 + count * 0.02)
        got = [f[2] for f, _ in frames]
        in_order = got == sorted(got) and all(f[1] == 0x81 for f, _ in frames)
        if must_all:
            c.check(f"{count} PING back to back", f"{count} frames", f"{count} responses in order",
                    f"{len(got)} responses", got == list(range(count)))
        else:
            c.check(f"{count} PING back to back", f"{count} frames",
                    ">= 8 responses, in order, no corrupted frame",
                    f"{len(got)} responses: {got}", len(got) >= 8 and in_order,
                    "frames beyond the queue depth may be dropped (NFR-04)")
        time.sleep(0.2)


@testcase("TC-11", "Responses decoded while the board logs", "SW-REQ-021, 022, 029")
def tc_shared_uart(c):
    c.link.flush(0.05)
    bad = 0
    for i in range(30):
        rsp, _, texts = c.link.transact(CMD_PING, i)
        if rsp != bytes([SOF, 0x81, i]):
            bad += 1
        if any(ord(ch) >= 0x80 for t in texts for ch in t):
            bad += 1
    c.check("30 PING with trace lines", "30 frames", "30 clean responses, ASCII-only log",
            f"{30 - bad} clean", bad == 0)


@testcase("TC-13", "Response time (host round trip)", "SW-REQ-019")
def tc_timing(c):
    c.link.flush(0.05)
    rtts = []
    for i in range(50):
        rsp, rtt, _ = c.link.transact(CMD_PING, i)
        if rtt is not None and rsp == bytes([SOF, 0x81, i]):
            rtts.append(rtt * 1000)
    if not rtts:
        c.check("Round-trip time", "50 x PING", "responses", "none", False)
        return
    avg, mx = sum(rtts) / len(rtts), max(rtts)
    c.check("Round-trip time", "50 x PING",
            f"max < {c.args.rt_limit:.0f} ms (incl. USB latency)",
            f"min {min(rtts):.1f} / avg {avg:.1f} / max {mx:.1f} ms",
            len(rtts) == 50 and mx < c.args.rt_limit,
            "board-side 20 ms is measured with a logic analyser")


# ------------------------------------------------------------- runner ---
def run(link, args):
    results = []
    selected = [t for t in TESTS if not args.only or t["id"] in args.only]
    for t in selected:
        print(f"\n{t['id']}  {t['title']}")
        ctx = Ctx(link, args)
        start = time.perf_counter()
        error = ""
        try:
            link.flush(0.1)
            t["fn"](ctx)
        except Exception as exc:                  # a crash in a test is a FAIL, not an abort
            error = f"{type(exc).__name__}: {exc}"
            ctx.check("Exception", "-", "-", error, False)
        for s in ctx.steps:
            mark = "\x1b[32mPASS\x1b[0m" if s["ok"] else "\x1b[31mFAIL\x1b[0m"
            print(f"   [{mark}] {s['step']:<32} sent {s['sent']:<18} exp {s['expected']:<22} got {s['received']}"
                  + (f"  ({s['note']})" if s["note"] else ""))
        verdict = "Passed" if ctx.steps and all(s["ok"] for s in ctx.steps) else "Failed"
        results.append({**t, "steps": ctx.steps, "verdict": verdict,
                        "duration": time.perf_counter() - start})
        print(f"   => {verdict}")
    return results


def write_reports(results, args):
    os.makedirs(args.out, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.join(args.out, f"apk5_report_{stamp}")
    passed = sum(r["verdict"] == "Passed" for r in results)
    target = "simulator" if args.sim else args.port

    with open(base + ".csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Test ID", "Title", "Verifies", "Step", "Sent", "Expected", "Received", "Step result", "Note", "Test verdict"])
        for r in results:
            for s in r["steps"]:
                w.writerow([r["id"], r["title"], r["verifies"], s["step"], s["sent"], s["expected"],
                            s["received"], "PASS" if s["ok"] else "FAIL", s["note"], r["verdict"]])

    rows = []
    for r in results:
        cls = "pass" if r["verdict"] == "Passed" else "fail"
        steps = "".join(
            f"<tr class='{'ok' if s['ok'] else 'ko'}'><td>{html.escape(s['step'])}</td><td><code>{html.escape(s['sent'])}</code></td>"
            f"<td><code>{html.escape(s['expected'])}</code></td><td><code>{html.escape(s['received'])}</code></td>"
            f"<td>{'PASS' if s['ok'] else 'FAIL'}</td><td>{html.escape(s['note'])}</td></tr>"
            for s in r["steps"])
        rows.append(
            f"<details {'open' if cls == 'fail' else ''}><summary><span class='badge {cls}'>{r['verdict'].upper()}</span> "
            f"<b>{r['id']}</b> {html.escape(r['title'])} <span class='muted'>· verifies {html.escape(r['verifies'])} · "
            f"{r['duration']:.2f} s</span></summary><table><tr><th>Step</th><th>Sent</th><th>Expected</th>"
            f"<th>Received</th><th>Result</th><th>Note</th></tr>{steps}</table></details>")
    page = f"""<!doctype html><html><head><meta charset="utf-8"><title>APK5 test report</title><style>
body{{font-family:Arial,sans-serif;margin:24px;color:#222}} h1{{color:#1F3864;margin-bottom:4px}}
.muted{{color:#666;font-weight:normal;font-size:13px}} table{{border-collapse:collapse;width:100%;margin:8px 0 16px;font-size:13px}}
th,td{{border:1px solid #ccc;padding:4px 8px;text-align:left}} th{{background:#1F3864;color:#fff}}
tr.ko td{{background:#FFC7CE}} summary{{cursor:pointer;padding:6px 0;font-size:15px}}
.badge{{display:inline-block;padding:2px 8px;border-radius:4px;color:#fff;font-size:12px}}
.pass{{background:#2E7D32}} .fail{{background:#C62828}} .sum td{{font-size:14px}}</style></head><body>
<h1>APK5 UART protocol - test report</h1>
<div class="muted">{datetime.datetime.now():%Y-%m-%d %H:%M:%S} · target {html.escape(str(target))} · tester {html.escape(args.tester)}</div>
<table class="sum" style="width:auto"><tr><th>Test cases</th><th>Passed</th><th>Failed</th><th>Verdict</th></tr>
<tr><td>{len(results)}</td><td>{passed}</td><td>{len(results) - passed}</td>
<td><span class="badge {'pass' if passed == len(results) else 'fail'}">{'PASS' if passed == len(results) else 'FAIL'}</span></td></tr></table>
{''.join(rows)}</body></html>"""
    with open(base + ".html", "w", encoding="utf-8") as f:
        f.write(page)
    return base


def update_xlsx(results, args):
    from openpyxl import load_workbook
    wb = load_workbook(args.xlsx)
    ws = wb["Test Specification"]
    today = datetime.date.today().isoformat()
    by_id = {r["id"]: r for r in results}
    for row in range(5, ws.max_row + 1):
        r = by_id.get(ws.cell(row=row, column=1).value)
        if r:
            ws.cell(row=row, column=8, value=r["verdict"])
            ws.cell(row=row, column=9, value=today)
            ws.cell(row=row, column=10, value=args.tester)
            failed = [s["step"] for s in r["steps"] if not s["ok"]]
            ws.cell(row=row, column=11, value=("Failed steps: " + ", ".join(failed)) if failed else "Automated run")
    wb.save(args.xlsx)


def main():
    ap = argparse.ArgumentParser(description="APK5 automated protocol tests")
    ap.add_argument("--port", help="COM port, e.g. COM5")
    ap.add_argument("--baud", type=int, default=115200)
    ap.add_argument("--sim", action="store_true", help="use the built-in board simulator")
    ap.add_argument("--only", nargs="*", help="test IDs to run, e.g. TC-01 TC-08")
    ap.add_argument("--interactive", action="store_true", help="run the steps that need an operator")
    ap.add_argument("--led-delay", type=float, default=0.05, help="pause between LED steps (s), raise to watch the LEDs")
    ap.add_argument("--rt-limit", type=float, default=30.0, help="max host round-trip time in ms for TC-13")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports"))
    ap.add_argument("--xlsx", help="workbook whose Test Specification sheet receives the results")
    ap.add_argument("--tester", default=os.environ.get("USERNAME", "tester"))
    args = ap.parse_args()

    if args.sim:
        port = SimBoard()
    elif args.port:
        import serial
        port = serial.Serial(args.port, args.baud, timeout=0.01)
    else:
        ap.error("give --port COMx or --sim")

    if os.name == "nt":
        os.system("")                              # enable ANSI colours in the Windows console
    link = Link(port)
    try:
        results = run(link, args)
    finally:
        link.close()

    base = write_reports(results, args)
    passed = sum(r["verdict"] == "Passed" for r in results)
    print(f"\n{passed}/{len(results)} test cases passed")
    print(f"Report: {base}.html\n        {base}.csv")
    if args.xlsx:
        update_xlsx(results, args)
        print(f"Results written to {args.xlsx}")
    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
